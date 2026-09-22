import os
import time
from pathlib import Path

import requests
from dotenv import load_dotenv

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

import pymupdf
try:
    import pytesseract
    from PIL import Image
except ImportError:
    pytesseract = None
    Image = None
from docx import Document


if pytesseract is not None:
    pytesseract.pytesseract.tesseract_cmd = (
        r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    )

def extract_elements_from_unstructured(file_path: str) -> list[dict]:
    api_key = os.getenv("UNSTRUCTURED_API_KEY")
    api_url = os.getenv(
        "UNSTRUCTURED_API_URL",
        "https://platform-api.transform.unstructured.io/api/v1"
    )

    if not api_key:
        raise ValueError(
            "UNSTRUCTURED_API_KEY is missing from the .env file."
        )

    content_type = {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }.get(Path(file_path).suffix.lower(), "application/octet-stream")

    if "platform-api.transform.unstructured.io" in api_url:
        return _extract_elements_from_platform_job(
            file_path, api_url, api_key, content_type
        )

    with open(file_path, "rb") as file:
        response = requests.post(
            api_url,
            headers={
                "unstructured-api-key": api_key
            },
            files={
                "files": (
                    os.path.basename(file_path),
                    file,
                    content_type
                )
            },
            data={
                "strategy": "hi_res",
                "languages": "eng",
            },
            timeout=120
        )

    try:
        response.raise_for_status()
    except requests.HTTPError as error:
        detail = response.text[:500].replace("\n", " ").strip()
        raise RuntimeError(
            f"Unstructured API request failed ({response.status_code}): {detail}"
        ) from error

    try:
        elements = response.json()
    except ValueError as error:
        raise RuntimeError("Unstructured API returned invalid JSON.") from error

    if not isinstance(elements, list):
        raise RuntimeError("Unstructured API returned an unexpected response shape.")
    return elements


def _extract_elements_from_platform_job(
    file_path: str,
    api_url: str,
    api_key: str,
    content_type: str,
) -> list[dict]:
    base_url = api_url.rstrip("/")
    headers = {"unstructured-api-key": api_key, "accept": "application/json"}
    with open(file_path, "rb") as file:
        response = requests.post(
            f"{base_url}/jobs/",
            headers=headers,
            files={
                "input_files": (
                    os.path.basename(file_path),
                    file,
                    content_type,
                )
            },
            data={
                "request_data": (
                    '{"template_id":"hi_res_partition"}'
                )
            },
            timeout=120,
        )
    try:
        response.raise_for_status()
        job = response.json()
    except (requests.HTTPError, ValueError) as error:
        detail = response.text[:500].replace("\n", " ").strip()
        raise RuntimeError(
            f"Unstructured job creation failed ({response.status_code}): {detail}"
        ) from error

    job_id = job.get("id") or job.get("job_id")
    if not job_id:
        raise RuntimeError("Unstructured job creation returned no job ID.")

    deadline = time.monotonic() + float(os.getenv("UNSTRUCTURED_TIMEOUT", "300"))
    while time.monotonic() < deadline:
        time.sleep(2)
        status_response = requests.get(
            f"{base_url}/jobs/{job_id}",
            headers=headers,
            timeout=30,
        )
        if not status_response.ok:
            continue
        status_data = status_response.json()
        status = str(status_data.get("status", "")).upper()
        if status in {"SCHEDULED", "IN_PROGRESS", "NEW", "PENDING", "PROCESSING", "RUNNING"}:
            continue
        if status not in {"COMPLETED", "FINISHED", "SUCCESS", "DONE", "COMPLETE"}:
            raise RuntimeError(f"Unstructured job {job_id} failed with status {status}.")
        return _download_platform_job_elements(
            base_url, job_id, headers, status_data
        )

    raise RuntimeError(
        f"Unstructured job {job_id} did not complete within the configured timeout."
    )


def _download_platform_job_elements(
    base_url: str,
    job_id: str,
    headers: dict[str, str],
    job_data: dict,
) -> list[dict]:
    details_response = requests.get(
        f"{base_url}/jobs/{job_id}/details",
        headers=headers,
        timeout=30,
    )
    details = details_response.json() if details_response.ok else {}
    file_ids: set[str] = set()

    def collect_file_ids(value: object) -> None:
        if isinstance(value, list):
            for item in value:
                collect_file_ids(item)
        elif isinstance(value, dict):
            file_id = value.get("file_id")
            if isinstance(file_id, str):
                file_ids.add(file_id)
            for item in value.values():
                collect_file_ids(item)

    collect_file_ids(job_data.get("output_node_files"))
    collect_file_ids(details.get("output_node_files"))
    collect_file_ids(details.get("node_file_metadata"))
    elements = []
    for file_id in file_ids:
        response = requests.get(
            f"{base_url}/jobs/{job_id}/download",
            params={"file_id": file_id},
            headers=headers,
            timeout=60,
        )
        if not response.ok:
            continue
        payload = response.json()
        if isinstance(payload, list):
            elements.extend(
                item for item in payload
                if isinstance(item, dict)
            )
        elif isinstance(payload, dict):
            for value in payload.values():
                if isinstance(value, list):
                    elements.extend(
                        item for item in value
                        if isinstance(item, dict)
                    )
    if not elements:
        raise RuntimeError("Unstructured job completed without readable output elements.")
    return elements


def _elements_to_text(elements: list[dict]) -> str:
    text_parts = []
    for element in elements:
        if not isinstance(element, dict):
            continue
        text = element.get("text")
        if isinstance(text, str) and text.strip():
            text_parts.append(text.strip())
    return "\n".join(text_parts).strip()

def extract_text_from_pdf_local(file_path: str) -> str:
    document = pymupdf.open(file_path)
    all_pages_text = []

    for page_number, page in enumerate(document):
        page_width = page.rect.width
        page_height = page.rect.height
        middle_x = page_width / 2

        blocks = page.get_text("blocks")

        left_blocks = []
        right_blocks = []
        full_width_blocks = []

        for block in blocks:
            x0, y0, x1, y1, text, *_ = block

            text = text.strip()

            if not text:
                continue

            block_width = x1 - x0

            # Blocks covering most of the page are probably
            # full-width headings or headers.
            if block_width > page_width * 0.75:
                full_width_blocks.append((y0, x0, text))

            # Left column
            elif x0 < middle_x:
                left_blocks.append((y0, x0, text))

            # Right column
            else:
                right_blocks.append((y0, x0, text))

        # Sort each group from top to bottom.
        # x0 is used when two blocks have the same y position.
        full_width_blocks.sort(key=lambda item: (item[0], item[1]))
        left_blocks.sort(key=lambda item: (item[0], item[1]))
        right_blocks.sort(key=lambda item: (item[0], item[1]))

        page_text = []

        # Full-width content first, such as name/contact information
        for _, _, text in full_width_blocks:
            page_text.append(text)

        # Read the complete left column
        for _, _, text in left_blocks:
            page_text.append(text)

        # Then read the complete right column
        for _, _, text in right_blocks:
            page_text.append(text)

        all_pages_text.append("\n".join(page_text))

    document.close()

    normal_text = "\n\n".join(all_pages_text).strip()

    # Use normal extraction if enough text is available.
    if len(normal_text) >= 50:
        return normal_text

    # Not enough text — this is a scanned PDF. Signal the caller to use vision.
    raise ScannedPDFError(file_path)


class ScannedPDFError(Exception):
    """Raised when a PDF yields too little text for LLM parsing (likely scanned)."""
    def __init__(self, file_path: str):
        super().__init__(f"PDF appears to be scanned or image-based: {file_path}")
        self.file_path = file_path


def pdf_to_page_images(file_path: str, max_pages: int = 5, scale: float = 2.0) -> list[bytes]:
    """Render each page of a PDF to a JPEG (bytes) using pymupdf. No Tesseract needed."""
    document = pymupdf.open(file_path)
    images: list[bytes] = []
    for page in document:
        if len(images) >= max_pages:
            break
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale))
        images.append(pixmap.tobytes("jpeg"))
    document.close()
    return images

def extract_text_from_docx_local(file_path: str) -> str:
    document = Document(file_path)

    extracted_text = []

    # Extract normal paragraphs.
    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            extracted_text.append(paragraph_text)

    # Extract tables.
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                extracted_text.append(" | ".join(row_text))

    return "\n".join(extracted_text).strip()


def extract_text(file_path: str) -> str:
    path = Path(file_path)
    extension = path.suffix.lower()

    # LM Studio now handles structured resume parsing. Unstructured remains an
    # opt-in text/OCR provider for deployments that explicitly enable it.
    if (
        extension in {".pdf", ".docx"}
        and os.getenv("UNSTRUCTURED_ENABLED", "false").lower() in {"1", "true", "yes"}
        and os.getenv("UNSTRUCTURED_API_KEY")
    ):
        try:
            extracted_text = _elements_to_text(
                extract_elements_from_unstructured(file_path)
            )
            if extracted_text:
                return extracted_text
            raise RuntimeError("Unstructured API returned no readable text.")
        except (requests.RequestException, RuntimeError) as error:
            if os.getenv("UNSTRUCTURED_FALLBACK_LOCAL", "true").lower() not in {
                "1", "true", "yes"
            }:
                raise ValueError(str(error)) from error

    if extension == ".pdf":
        return extract_text_from_pdf_local(file_path)

    elif extension == ".docx":
        return extract_text_from_docx_local(file_path)

    elif extension == ".txt":
        return path.read_text(
            encoding="utf-8",
            errors="replace"
        ).strip()

    else:
        raise ValueError(
            "Unsupported file type. "
            "Only PDF, DOCX, and TXT files are supported."
        )